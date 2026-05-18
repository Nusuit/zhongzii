from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "assets" / "vocab_hsk1_6.xlsx"
OUTPUT = ROOT / "lib" / "data" / "hsk_vocab_data.dart"


def escape_dart_string(value: str) -> str:
    return (
        value.replace("\\", "\\\\")
        .replace("'", "\\'")
        .replace("\r", "")
        .replace("\n", "\\n")
    )


def to_text(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def main() -> None:
    wb = load_workbook(INPUT, read_only=True, data_only=True)
    rows: list[tuple[int, str, str, str, str]] = []

    for level in range(1, 7):
        sheet_name = f"HSK{level}"
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue

            hanzi = to_text(row[1] if len(row) > 1 else None)
            pinyin = to_text(row[2] if len(row) > 2 else None)
            meaning = to_text(row[3] if len(row) > 3 else None)
            ex_hanzi = to_text(row[4] if len(row) > 4 else None)
            ex_pinyin = to_text(row[5] if len(row) > 5 else None)
            ex_vi = to_text(row[6] if len(row) > 6 else None)

            if not hanzi or not meaning:
                continue

            example_parts = []
            if ex_hanzi:
                example_parts.append(ex_hanzi)
            if ex_pinyin:
                example_parts.append(f"Pinyin: {ex_pinyin}")
            if ex_vi:
                example_parts.append(f"Dich: {ex_vi}")

            rows.append((level, hanzi, pinyin, meaning, " | ".join(example_parts)))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", encoding="utf-8") as f:
        f.write("// GENERATED CODE - DO NOT EDIT.\n")
        f.write("// Source: assets/vocab_hsk1_6.xlsx\n\n")
        f.write("class HardcodedVocabRecord {\n")
        f.write("  const HardcodedVocabRecord({\n")
        f.write("    required this.level,\n")
        f.write("    required this.hanzi,\n")
        f.write("    required this.pinyin,\n")
        f.write("    required this.meaning,\n")
        f.write("    required this.example,\n")
        f.write("  });\n\n")
        f.write("  final int level;\n")
        f.write("  final String hanzi;\n")
        f.write("  final String pinyin;\n")
        f.write("  final String meaning;\n")
        f.write("  final String example;\n")
        f.write("}\n\n")
        f.write("const hardcodedVocabulary = <HardcodedVocabRecord>[\n")

        for level, hanzi, pinyin, meaning, example in rows:
            f.write("  HardcodedVocabRecord(\n")
            f.write(f"    level: {level},\n")
            f.write(f"    hanzi: '{escape_dart_string(hanzi)}',\n")
            f.write(f"    pinyin: '{escape_dart_string(pinyin)}',\n")
            f.write(f"    meaning: '{escape_dart_string(meaning)}',\n")
            f.write(f"    example: '{escape_dart_string(example)}',\n")
            f.write("  ),\n")

        f.write("];\n")

    print(f"Generated {len(rows)} records -> {OUTPUT}")


if __name__ == "__main__":
    main()
